#!/usr/bin/env python3
"""Generates reference/example-template.pptx: a small, generic style template
built from scratch (solid-color backgrounds, standard system fonts, a
placeholder logo shape) that demonstrates the 5 reusable slide "shapes"
documented in reference/slide-types.md. It does not reuse any real church's
background photography, custom fonts, or logo -- point the real tool at your
own church's most recent deck for actual styling.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

WIDTH = Emu(18288000)
HEIGHT = Emu(10287000)

NAVY = RGBColor(0x1B, 0x1F, 0x3B)
CREAM = RGBColor(0xF5, 0xEE, 0xDC)
GOLD = RGBColor(0xD9, 0xB9, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_background(slide, color):
    # A full-bleed rectangle shape (not the p:bg slide property) -- this is
    # how the real church decks do it too, and it renders consistently
    # across PowerPoint and LibreOffice.
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, HEIGHT)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def add_logo_placeholder(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(700000), Emu(9100000), Emu(500000), Emu(500000))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()
    tf = shape.text_frame
    tf.text = "LOGO"
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.color.rgb = NAVY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER


def add_textbox(slide, left, top, width, height, lines, size, color, bold=True, italic=False, align=PP_ALIGN.CENTER, font="Calibri"):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def build():
    prs = Presentation()
    prs.slide_width = WIDTH
    prs.slide_height = HEIGHT
    blank = prs.slide_layouts[6]

    # 1-2: blank "waiting" slides
    for _ in range(2):
        s = prs.slides.add_slide(blank)
        add_background(s, NAVY)

    # 3: purpose / mission statement slide (placeholder org text, not a real church's)
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_logo_placeholder(s)
    add_textbox(s, 5700000, 1200000, 6900000, 900000, ["declaracion de"], 26, GOLD, bold=False, italic=True)
    add_textbox(s, 5700000, 1900000, 6900000, 900000, ["proposito"], 44, WHITE, bold=True, italic=True)
    add_textbox(s, 3300000, 3300000, 11700000, 3000000,
                ["Texto de ejemplo para la declaracion de proposito de la organizacion.",
                 "Sustituye este parrafo con el texto real de tu iglesia."],
                28, WHITE)

    # 4: hymn title slide (title + "Himno ##")
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 1700000, 4500000, 14900000, 1600000, ["Nombre del Himno de Ejemplo"], 40, CREAM, italic=True)
    add_textbox(s, 1700000, 6200000, 14900000, 1600000, ["Himno 00"], 46, GOLD, italic=True)

    # 5: lyric slide (verse couplet)
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 1600000, 3000000, 15100000, 4300000,
                ["Linea de ejemplo uno de la letra,", "linea de ejemplo dos de la letra."],
                40, WHITE)

    # 6: contemporary song title slide (2-line stacked, no hymn number)
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 2000000, 3200000, 14300000, 1600000, ["Titulo de"], 40, WHITE, italic=True)
    add_textbox(s, 2000000, 4800000, 14300000, 2200000, ["Cancion Contemporanea"], 58, CREAM, italic=True)

    # 7: lyric slide (second example, longer stanza)
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 1600000, 2600000, 15100000, 5000000,
                ["Primera linea de ejemplo,", "segunda linea de ejemplo,", "tercera linea de ejemplo."],
                36, WHITE)

    # 8: scripture section title
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_logo_placeholder(s)
    add_textbox(s, 1600000, 2600000, 15100000, 900000, ["Subtitulo de ejemplo"], 26, CREAM, italic=True)
    add_textbox(s, 1600000, 3300000, 15100000, 1600000, ["Titulo de la Seccion"], 54, WHITE, bold=True)
    add_textbox(s, 1600000, 5200000, 15100000, 900000, ["Libro 00:0-0"], 30, GOLD, bold=True)

    # 9: scripture reference sub-slide
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 4300000, 3700000, 9700000, 900000, ["Libro"], 34, WHITE, italic=True)
    add_textbox(s, 4300000, 4700000, 9700000, 900000, ["00:0-0"], 34, WHITE, italic=True)

    # 10: scripture verse text slide
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, 1000000, 3200000, 16300000, 3800000,
                ["1", "Texto de ejemplo del versiculo uno.", "", "2", "Texto de ejemplo del versiculo dos."],
                28, WHITE, bold=False, align=PP_ALIGN.LEFT)

    prs.save("reference/example-template.pptx")
    print("wrote reference/example-template.pptx —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    build()
